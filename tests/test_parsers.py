import io

import openpyxl
import pytest

from app.parsers.xlsx_parser import XlsxParser
from app.parsers.factory import is_xlsx, get_mime_type

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_xlsx(sheet_name: str, rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_parser_extracts_sheet_name():
    content = _make_xlsx("Payable Ageing", [["Vendor", "Amount"]])
    text = XlsxParser().parse(content)
    assert "[Sheet: Payable Ageing]" in text


def test_xlsx_parser_extracts_cell_values():
    content = _make_xlsx("Sheet1", [["CR Number", "123456789"]])
    text = XlsxParser().parse(content)
    assert "CR Number" in text
    assert "123456789" in text


def test_xlsx_parser_skips_empty_rows():
    content = _make_xlsx("Sheet1", [["Value"], [], ["Another"]])
    text = XlsxParser().parse(content)
    assert "Value" in text
    assert "Another" in text


def test_xlsx_parser_multiple_sheets():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["col1"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["col2"])
    buf = io.BytesIO()
    wb.save(buf)
    text = XlsxParser().parse(buf.getvalue())
    assert "[Sheet: Sheet1]" in text
    assert "[Sheet: Sheet2]" in text


def test_is_xlsx_by_mime():
    assert is_xlsx(XLSX_MIME, "file.xlsx") is True


def test_is_xlsx_by_extension():
    assert is_xlsx("application/octet-stream", "report.xlsx") is True


def test_is_xlsx_false_for_pdf():
    assert is_xlsx("application/pdf", "doc.pdf") is False


def test_get_mime_type_pdf():
    assert get_mime_type("application/pdf", "doc.pdf") == "application/pdf"


def test_get_mime_type_fallback_by_extension():
    assert get_mime_type("application/octet-stream", "doc.pdf") == "application/pdf"


def test_get_mime_type_jpeg():
    assert get_mime_type("image/jpeg", "photo.jpg") == "image/jpeg"
